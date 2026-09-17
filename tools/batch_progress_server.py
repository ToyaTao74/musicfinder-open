#!/usr/bin/env python3
"""授权表批量查询 —— 实时进度页面服务。

用法：./venv/bin/python3 tools/batch_progress_server.py [端口]
浏览器打开 http://127.0.0.1:8766/ 即可看到实时进度（页面每 5 秒自动刷新）。
只读结果 xlsx，与跑批进程互不干扰；服务挂了刷新页面会失败，重启本脚本即可。
"""
import html
import os
import sys
import time
from http.server import BaseHTTPRequestHandler, HTTPServer

import openpyxl

XLSX = "/Users/toya/Desktop/宜昌汇报/0902 AI会/02-工具/授权表查询_三平台收藏量.xlsx"
PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 8766
TOTAL = 869


def stats():
    import glob
    files = [XLSX] + sorted(glob.glob(XLSX.replace('.xlsx', '_part*.xlsx')))
    if not any(os.path.exists(f) for f in files):
        return {'n': 0, 'started': False}
    rows_all = []
    qq = wy = kg = miss = doubt = 0
    last = '—'
    mtime = 0
    for f in files:
        if not os.path.exists(f):
            continue
        try:
            ws = openpyxl.load_workbook(f, data_only=True).active
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
        except Exception:
            continue
        rows_all += rows
        qq += sum(1 for r in rows if len(r) > 4 and r[4] and '精准' in str(r[4]))
        wy += sum(1 for r in rows if len(r) > 7 and r[7] and '精准' in str(r[7]))
        kg += sum(1 for r in rows if len(r) > 10 and r[10] and '精准' in str(r[10]))
        miss += sum(1 for r in rows
                    if (len(r) <= 4 or not r[4] or '未收录' in str(r[4]))
                    and (len(r) <= 7 or not r[7] or '未收录' in str(r[7]))
                    and (len(r) <= 10 or not r[10] or '未收录' in str(r[10])))
        doubt += sum(1 for r in rows
                     if any(len(r) > i and '存疑' in str(r[i]) for i in (4, 7, 10)))
        if rows:
            last = f"{rows[-1][0]} — {rows[-1][1] if len(rows[-1]) > 1 and rows[-1][1] else ''}"
        mtime = max(mtime, os.path.getmtime(f))
    if not rows_all:
        return {'n': 0, 'started': False}
    age = int(time.time() - mtime) if mtime else 9999
    return {'n': len(rows_all), 'started': True, 'qq': qq, 'wy': wy, 'kg': kg,
            'miss': miss, 'doubt': doubt, 'age': age, 'last': last,
            'mtime': time.strftime('%H:%M:%S', time.localtime(mtime)) if mtime else '—'}


def page():
    s = stats()
    n = s['n']
    pct = n / TOTAL * 100
    bar_n = int(pct / 4)
    bar = '█' * bar_n + '░' * (25 - bar_n)
    running = '🟢 正在查询' if s['age'] < 300 else '🔴 已停滞（进程可能停了）'
    recent = []
    if s['started']:
        try:
            ws = openpyxl.load_workbook(XLSX, data_only=True).active
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]][-5:]
            recent = [(str(r[0]), str(r[1]) if len(r) > 1 and r[1] else '') for r in rows]
        except Exception:
            pass
    recent_html = ''.join(
        f'<tr><td>{html.escape(a)}</td><td>{html.escape(b)}</td></tr>' for a, b in reversed(recent)
    ) or '<tr><td colspan="2">（首批落盘前无数据）</td></tr>'
    return f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8">
<meta http-equiv="refresh" content="5">
<title>授权表查询进度</title>
<style>
body{{font-family:-apple-system,"PingFang SC","Microsoft YaHei",sans-serif;background:linear-gradient(160deg,#eef4ff,#f6f4ff);min-height:100vh;display:flex;align-items:center;justify-content:center;margin:0}}
.card{{background:#fff;border-radius:20px;box-shadow:0 12px 40px rgba(50,70,120,.12);padding:34px 40px;width:520px}}
h1{{font-size:20px;color:#1c2433;margin:0 0 4px}}
.st{{font-size:13px;color:#8a93a6;margin-bottom:16px}}
.bar{{background:#eef1f6;border-radius:10px;height:26px;overflow:hidden}}
.fill{{background:linear-gradient(90deg,#4f6ef7,#7c5cf7);height:100%;border-radius:10px;transition:width .5s}}
.pct{{font-size:26px;font-weight:700;color:#1c2433;margin:14px 0 2px}}
.n{{font-size:13px;color:#6b7280}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin-top:18px}}
.item{{background:#f7f8fb;border-radius:12px;padding:12px 16px}}
.item b{{font-size:20px;color:#3b5bfd}}
.item span{{display:block;font-size:12px;color:#8a93a6;margin-top:2px}}
.warn{{color:#d97706;font-size:12px;margin-top:10px}}
table{{width:100%;border-collapse:collapse;margin-top:16px;font-size:13px}}
td{{padding:6px 8px;border-top:1px solid #f0f1f5;color:#4b5563}}
.foot{{font-size:11px;color:#b3bac6;margin-top:14px;text-align:center}}
</style></head><body><div class="card">
<h1>授权表 · 三平台收藏量查询</h1>
<div class="st">状态: {running} ｜ 数据更新于 {s['mtime']}（页面每 5 秒自动刷新）</div>
<div class="bar"><div class="fill" style="width:{pct:.1f}%"></div></div>
<div class="pct">{pct:.1f}%</div>
<div class="n">{n} / {TOTAL} 首　[{bar}]</div>
<div class="grid">
<div class="item"><b>{s.get('qq', 0)}</b><span>QQ音乐 精准命中</span></div>
<div class="item"><b>{s.get('wy', 0)}</b><span>网易云 精准命中</span></div>
<div class="item"><b>{s.get('kg', 0)}</b><span>酷狗 精准命中</span></div>
<div class="item"><b>{s.get('miss', 0)}</b><span>三平台全未收录（存疑 {s.get('doubt', 0)}）</span></div>
</div>
{s.get('last') and f'<div class="warn">最近落盘：{html.escape(str(s["last"]))}</div>' or ''}
<table><tr><td><b>最近 5 首</b></td><td></td></tr>{recent_html}</table>
<div class="foot">MusicFinder 批量查询 · 严格双要素精准匹配（含艺名变体）</div>
</div></body></html>"""


class H(BaseHTTPRequestHandler):
    def do_GET(self):
        body = page().encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *a):
        pass


if __name__ == '__main__':
    print(f'进度页: http://127.0.0.1:{PORT}/  （Ctrl+C 停止）')
    HTTPServer(('127.0.0.1', PORT), H).serve_forever()
